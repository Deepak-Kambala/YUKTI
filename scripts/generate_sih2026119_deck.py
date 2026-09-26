from __future__ import annotations

import csv
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "SIH2026119.pptx"

W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(31, 73, 125)
ORANGE = RGBColor(255, 110, 31)
GREEN = RGBColor(20, 140, 65)
BLACK = RGBColor(0, 0, 0)
DARK = RGBColor(46, 52, 64)
GREY = RGBColor(238, 241, 245)
MID = RGBColor(105, 116, 128)
LIGHT_BLUE = RGBColor(228, 238, 250)
LIGHT_GREEN = RGBColor(232, 245, 237)
LIGHT_ORANGE = RGBColor(255, 239, 229)


def rgb(hexstr: str) -> RGBColor:
    hexstr = hexstr.strip("#")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


def add_text(slide, text, x, y, w, h, size=18, color=BLACK, bold=False, font="Arial", align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    run = p.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line if line else fill
    shp.line.width = Pt(1)
    return shp


def add_title(slide, title, num):
    add_text(slide, "SMART INDIA HACKATHON 2026", Inches(0.75), Inches(0.18), Inches(8.4), Inches(0.45),
             29, NAVY, True, "Georgia")
    add_text(slide, "PS SIH26119", Inches(10.25), Inches(0.22), Inches(2.15), Inches(0.3),
             14, NAVY, True, "Arial", PP_ALIGN.RIGHT)
    add_text(slide, title, Inches(0.75), Inches(0.82), Inches(11.9), Inches(0.48),
             24, BLACK, True, "Georgia", PP_ALIGN.CENTER)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.75), Inches(1.38), Inches(11.85), Inches(0.035))
    line.fill.solid()
    line.fill.fore_color.rgb = NAVY
    line.line.color.rgb = NAVY
    add_text(slide, f"{num:02d}", Inches(12.0), Inches(6.95), Inches(0.45), Inches(0.22), 9, MID, False)


def add_chip(slide, text, x, y, fill, color=DARK, width=1.8):
    add_rect(slide, x, y, Inches(width), Inches(0.34), fill, fill, True)
    add_text(slide, text, x + Inches(0.08), y + Inches(0.06), Inches(width - 0.16), Inches(0.18), 9, color, True)


def add_bullet(slide, text, x, y, w, size=20, color=BLACK, accent=None):
    add_text(slide, "•", x, y, Inches(0.24), Inches(0.28), size + 6, accent or BLACK, True)
    return add_text(slide, text, x + Inches(0.32), y + Inches(0.05), w - Inches(0.32), Inches(0.42), size, color, True)


def add_card(slide, title, body, x, y, w, h, fill=RGBColor(255, 255, 255), border=RGBColor(214, 221, 230), title_color=NAVY):
    add_rect(slide, x, y, w, h, fill, border, True)
    add_text(slide, title, x + Inches(0.18), y + Inches(0.14), w - Inches(0.36), Inches(0.28), 15, title_color, True)
    add_text(slide, body, x + Inches(0.18), y + Inches(0.52), w - Inches(0.36), h - Inches(0.62), 12, DARK)


def arrow(slide, x1, y1, x2, y2, color=NAVY, width=2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    c.line.end_arrowhead = True
    return c


def load_benchmark():
    with open(ROOT / "results" / "benchmark_results.csv", newline="") as f:
        return list(csv.DictReader(f))


def load_scaling():
    with open(ROOT / "results" / "summaries" / "scaling_ab.csv", newline="") as f:
        return list(csv.DictReader(f))


def clean_runtime(v):
    try:
        return float(v)
    except Exception:
        return 0.0


def small_bar(slide, labels, values, x, y, w, h, colors, max_label=None):
    maxv = max(values) if values else 1
    bar_h = h / len(values) * 0.44
    gap = h / len(values) * 0.56
    for i, (lab, val) in enumerate(zip(labels, values)):
        yy = y + int(i * (bar_h + gap))
        add_text(slide, lab, x, yy - Inches(0.02), Inches(2.4), Inches(0.22), 10, DARK, True)
        add_rect(slide, x + Inches(2.55), yy, int((w - Inches(3.25)) * (val / maxv)), int(bar_h), colors[i], colors[i], True)
        add_text(slide, f"{val:g}{max_label or ''}", x + w - Inches(0.62), yy - Inches(0.02), Inches(0.7), Inches(0.22), 10, DARK, True)


def add_watermark(slide):
    # Light technical watermark, inspired by the supplied SIH reference but without unverified logo assets.
    for i, (x, y, s) in enumerate([(8.4, 1.35, 2.6), (7.2, 4.8, 1.0), (10.5, 4.6, 1.6)]):
        shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.HEXAGON, Inches(x), Inches(y), Inches(s), Inches(s * 0.86))
        shp.fill.solid()
        shp.fill.fore_color.rgb = GREY
        shp.fill.transparency = 38 if i == 0 else 55
        shp.line.color.rgb = GREY
        shp.line.transparency = 70


def slide_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    add_watermark(slide)
    add_text(slide, "SMART INDIA HACKATHON 2026", Inches(1.5), Inches(0.25), Inches(8.8), Inches(0.58), 34, NAVY, True, "Georgia")
    add_text(slide, "TITLE PAGE", Inches(4.75), Inches(1.2), Inches(3.9), Inches(0.52), 28, BLACK, True, "Georgia", PP_ALIGN.CENTER)
    add_text(slide, "Indigenous GPU-Accelerated Optimization Solver", Inches(1.02), Inches(1.95), Inches(11.2), Inches(0.42), 25, ORANGE, True)
    add_text(slide, "A verified from-scratch LP/MILP prototype today; a GPU/C++17 industrial solver roadmap next.", Inches(1.1), Inches(2.42), Inches(10.7), Inches(0.32), 16, GREEN, True, align=PP_ALIGN.CENTER)

    rows = [
        ("Problem Statement ID", "SIH26119", 20),
        ("Problem Statement Title", "Indigenous GPU-Accelerated Optimization Solver", 18),
        ("Theme", "Smart Automation", 20),
        ("PS Category", "Software", 20),
        ("Prototype", "Sovereign Optimization Engine", 20),
    ]
    y = Inches(3.02)
    for label, val, size in rows:
        add_bullet(slide, f"{label} -", Inches(0.85), y, Inches(4.8), 22)
        add_text(slide, val, Inches(5.1), y + Inches(0.05), Inches(7.6), Inches(0.38), size, NAVY, True)
        y += Inches(0.72)

    add_rect(slide, Inches(10.55), Inches(6.13), Inches(0.86), Inches(0.62), ORANGE, ORANGE)
    add_rect(slide, Inches(11.41), Inches(6.13), Inches(0.86), Inches(0.62), rgb("00689D"), rgb("00689D"))
    add_rect(slide, Inches(12.27), Inches(6.13), Inches(0.86), Inches(0.62), rgb("BF8B2E"), rgb("BF8B2E"))
    add_text(slide, "9", Inches(10.62), Inches(6.23), Inches(0.36), Inches(0.18), 16, RGBColor(255, 255, 255), True)
    add_text(slide, "16", Inches(11.47), Inches(6.23), Inches(0.42), Inches(0.18), 16, RGBColor(255, 255, 255), True)
    add_text(slide, "12", Inches(12.33), Inches(6.23), Inches(0.42), Inches(0.18), 16, RGBColor(255, 255, 255), True)


def slide_proposed(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_watermark(slide)
    add_title(slide, "PROPOSED SOLUTION", 2)
    add_text(slide, "Sovereign Optimization Engine", Inches(0.88), Inches(1.75), Inches(4.2), Inches(0.38), 24, NAVY, True)
    add_text(slide, "A transparent solver core for Indian industrial LP/MILP workloads, built without external optimization-solver libraries.", Inches(0.9), Inches(2.22), Inches(4.45), Inches(0.62), 16, DARK)

    x0, y0 = Inches(5.8), Inches(1.72)
    stages = [
        ("Model", "Variables\nconstraints\nobjective", LIGHT_BLUE),
        ("Presolve", "Fixed vars\nempty rows\nsingleton bounds", LIGHT_GREEN),
        ("LP Core", "Two-phase\nrevised simplex", LIGHT_ORANGE),
        ("MILP", "Branch-and-\nbound", LIGHT_BLUE),
        ("Verify", "Independent\nre-check", LIGHT_GREEN),
    ]
    x = x0
    for i, (t, b, c) in enumerate(stages):
        add_card(slide, t, b, x, y0, Inches(1.28), Inches(1.12), c, rgb("C8D2DF"))
        if i < len(stages) - 1:
            arrow(slide, x + Inches(1.28), y0 + Inches(0.56), x + Inches(1.58), y0 + Inches(0.56), ORANGE, 2)
        x += Inches(1.68)

    add_text(slide, "VALIDATED NOW", Inches(0.85), Inches(3.58), Inches(2.2), Inches(0.3), 16, GREEN, True)
    add_text(slide, "BUILD NEXT", Inches(7.05), Inches(3.58), Inches(2.0), Inches(0.3), 16, ORANGE, True)
    now = [
        ("LP/MILP core", "Two-phase revised simplex + branch-and-bound."),
        ("No solver dependency", "No CPLEX/Gurobi/HiGHS/CBC/GLPK/OR-Tools/PuLP/Pyomo/CVXPY/scipy.optimize import."),
        ("Fail-loud correctness", "OPTIMAL implies independent verification passed."),
    ]
    nxt = [
        ("C++17 port", "Port validated algorithms to the PS-preferred production stack."),
        ("Sparse LU", "Replace dense O(m^3) basis solves and add updates."),
        ("GPU investigation", "Profile first; accelerate only kernels with measured benefit."),
    ]
    yy = Inches(4.02)
    for t, b in now:
        add_card(slide, t, b, Inches(0.82), yy, Inches(5.45), Inches(0.65), RGBColor(255, 255, 255), rgb("C7E3D0"), GREEN)
        yy += Inches(0.82)
    yy = Inches(4.02)
    for t, b in nxt:
        add_card(slide, t, b, Inches(6.95), yy, Inches(5.55), Inches(0.65), RGBColor(255, 255, 255), rgb("F2D1BD"), ORANGE)
        yy += Inches(0.82)


def slide_how(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "HOW IT ADDRESSES THE PROBLEM", 3)
    add_text(slide, "BEFORE", Inches(0.9), Inches(1.75), Inches(2.0), Inches(0.34), 18, ORANGE, True)
    add_text(slide, "AFTER", Inches(9.9), Inches(1.75), Inches(2.0), Inches(0.34), 18, GREEN, True)

    add_card(slide, "Opaque foreign solver dependency", "Industrial planning depends on commercial black-box engines.", Inches(0.85), Inches(2.25), Inches(3.35), Inches(0.95), LIGHT_ORANGE, rgb("F2D1BD"), ORANGE)
    add_card(slide, "Unverifiable result path", "A status field can be trusted only if feasibility and objective are rechecked.", Inches(0.85), Inches(3.42), Inches(3.35), Inches(0.95), LIGHT_ORANGE, rgb("F2D1BD"), ORANGE)
    add_card(slide, "Scaling risk hidden", "Ill-conditioned models can fail or silently mislead if not tested.", Inches(0.85), Inches(4.59), Inches(3.35), Inches(0.95), LIGHT_ORANGE, rgb("F2D1BD"), ORANGE)

    for yy in [2.72, 3.89, 5.06]:
        arrow(slide, Inches(4.45), Inches(yy), Inches(8.55), Inches(yy), NAVY, 2.6)
        add_text(slide, "validated prototype", Inches(5.45), Inches(yy - 0.22), Inches(2.1), Inches(0.22), 11, NAVY, True, align=PP_ALIGN.CENTER)

    add_card(slide, "From-scratch solver core", "Basis solves, simplex pivots and B&B search live in this repo.", Inches(8.9), Inches(2.25), Inches(3.35), Inches(0.95), LIGHT_GREEN, rgb("C7E3D0"), GREEN)
    add_card(slide, "Independent verification", "Every reported solution is recomputed against the original model.", Inches(8.9), Inches(3.42), Inches(3.35), Inches(0.95), LIGHT_GREEN, rgb("C7E3D0"), GREEN)
    add_card(slide, "Measured numerical guard", "Threshold-gated equilibration solves ill-conditioned test cases.", Inches(8.9), Inches(4.59), Inches(3.35), Inches(0.95), LIGHT_GREEN, rgb("C7E3D0"), GREEN)

    add_text(slide, "Impact discipline: no external-solver speed claims, no commercial comparison, no fake benchmark data.", Inches(1.35), Inches(6.28), Inches(10.5), Inches(0.36), 15, NAVY, True, align=PP_ALIGN.CENTER)


def slide_innovation(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "INNOVATION AND UNIQUENESS", 4)
    add_text(slide, "What is unique in this SIH prototype", Inches(0.85), Inches(1.65), Inches(5.5), Inches(0.32), 18, NAVY, True)

    cols = [
        ("Sovereign core", "Optimization algorithms are implemented in project code, not delegated to solver libraries.", LIGHT_BLUE, NAVY),
        ("Fail-loud verification", "A failed independent recheck downgrades status to NUMERICAL_FAILURE.", LIGHT_GREEN, GREEN),
        ("Measured scaling fix", "Power-of-two equilibration repairs two ill-conditioned LP failures at 1.00x same-process runtime.", LIGHT_ORANGE, ORANGE),
    ]
    x = Inches(0.82)
    for t, b, f, c in cols:
        add_card(slide, t, b, x, Inches(2.05), Inches(3.88), Inches(1.35), f, rgb("D4DAE4"), c)
        x += Inches(4.18)

    add_text(slide, "Equilibration A/B outcome", Inches(0.85), Inches(3.95), Inches(3.5), Inches(0.3), 16, BLACK, True)
    small_bar(slide, ["Scaling off", "Auto default", "Always on"], [13, 15, 15],
              Inches(0.95), Inches(4.43), Inches(5.5), Inches(1.4), [ORANGE, GREEN, NAVY], "/15 solved")
    small_bar(slide, ["Known optima off", "Known optima auto"], [6, 8],
              Inches(7.0), Inches(4.5), Inches(5.0), Inches(1.0), [ORANGE, GREEN], "/8")
    add_text(slide, "Build next: sparse factorization, warm starts, C++17, reference benchmarks, and GPU only after profiling.", Inches(1.15), Inches(6.35), Inches(10.8), Inches(0.32), 15, DARK, True, align=PP_ALIGN.CENTER)


def slide_technical(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "TECHNICAL APPROACH", 5)
    y = Inches(1.85)
    stages = [
        ("Problem representation", "Variable, Constraint, Model"),
        ("Presolve", "fixed variables, empty rows, singleton bound tightening"),
        ("LP solver", "two-phase revised simplex; Bland's rule"),
        ("Linear algebra", "own Gaussian elimination with partial pivoting"),
        ("MILP layer", "LP relaxations inside branch-and-bound"),
        ("Verification", "bounds, constraints, integrality, objective recomputed"),
    ]
    x = Inches(0.9)
    for i, (t, b) in enumerate(stages):
        add_card(slide, t, b, x, y, Inches(1.72), Inches(1.08), RGBColor(255, 255, 255), rgb("C8D2DF"), NAVY if i != 5 else GREEN)
        if i < len(stages) - 1:
            arrow(slide, x + Inches(1.72), y + Inches(0.54), x + Inches(2.02), y + Inches(0.54), ORANGE, 2)
        x += Inches(2.02)

    add_text(slide, "Repository module map", Inches(0.95), Inches(3.5), Inches(2.6), Inches(0.28), 15, BLACK, True)
    modules = [
        ("core", "model representation"),
        ("linalg", "Gaussian solve + scaling"),
        ("lp", "simplex + orchestrator"),
        ("milp", "branch-and-bound"),
        ("io", "basic MPS I/O"),
        ("industrial", "synthetic refinery cases"),
    ]
    x, y = Inches(0.95), Inches(3.98)
    for i, (m, desc) in enumerate(modules):
        add_chip(slide, m, x, y, LIGHT_BLUE, NAVY, 1.25)
        add_text(slide, desc, x + Inches(1.38), y + Inches(0.06), Inches(2.35), Inches(0.2), 10, DARK)
        if i % 2 == 1:
            x = Inches(0.95)
            y += Inches(0.55)
        else:
            x += Inches(5.95)

    add_text(slide, "Current prototype scope boundary", Inches(0.95), Inches(5.85), Inches(3.2), Inches(0.28), 15, BLACK, True)
    add_text(slide, "Validated Python algorithm prototype; not yet sparse-LU, C++17, GPU, dual simplex, cuts, or production-scale.", Inches(0.95), Inches(6.25), Inches(11.2), Inches(0.28), 14, ORANGE, True)


def slide_feasibility(prs, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "FEASIBILITY AND VIABILITY", 6)
    optimal = sum(1 for r in rows if r["status"] == "OPTIMAL")
    infeas = sum(1 for r in rows if r["status"] == "INFEASIBLE")
    verified = sum(1 for r in rows if r["verified_feasible"] == "True")
    known = sum(1 for r in rows if r["known_optimal"])
    matched = sum(1 for r in rows if r["matches_known"] == "True")
    avg = sum(clean_runtime(r["runtime_sec"]) for r in rows) / len(rows)

    stats = [("15", "instances run"), (str(optimal), "optimal"), (str(infeas), "infeasible detected"), (f"{verified}/15", "verified"), (f"{matched}/{known}", "known optima matched"), (f"{avg:.4f}s", "avg runtime")]
    x = Inches(0.8)
    for val, lab in stats:
        add_rect(slide, x, Inches(1.78), Inches(1.85), Inches(1.05), RGBColor(255, 255, 255), rgb("C8D2DF"), True)
        add_text(slide, val, x, Inches(1.92), Inches(1.85), Inches(0.34), 24, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(slide, lab, x + Inches(0.1), Inches(2.35), Inches(1.65), Inches(0.26), 10, DARK, True, align=PP_ALIGN.CENTER)
        x += Inches(2.05)

    add_text(slide, "Runtime growth exposes the next engineering need", Inches(0.85), Inches(3.32), Inches(5.4), Inches(0.28), 15, BLACK, True)
    small_bar(slide, ["refinery small", "refinery medium", "refinery large"],
              [0.0028, 0.0854, 1.8945], Inches(0.9), Inches(3.78), Inches(5.65), Inches(1.45),
              [GREEN, NAVY, ORANGE], "s")

    add_text(slide, "Viability reading", Inches(7.0), Inches(3.32), Inches(2.6), Inches(0.28), 15, BLACK, True)
    add_card(slide, "Validated now", "Correctness and verification path are demonstrated on small benchmark and synthetic refinery cases.", Inches(7.0), Inches(3.75), Inches(5.2), Inches(0.82), LIGHT_GREEN, rgb("C7E3D0"), GREEN)
    add_card(slide, "Not claimed yet", "Large-scale performance, GPU speedup, Netlib/MIPLIB validation, and commercial-solver comparison are not established.", Inches(7.0), Inches(4.72), Inches(5.2), Inches(0.98), LIGHT_ORANGE, rgb("F2D1BD"), ORANGE)
    add_card(slide, "Next viability milestone", "Sparse LU with updates plus real MPS benchmark validation before any production claim.", Inches(7.0), Inches(5.85), Inches(5.2), Inches(0.72), LIGHT_BLUE, rgb("C8D2DF"), NAVY)


def slide_impact(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "IMPACT AND BENEFITS", 7)
    add_text(slide, "BEFORE -> AFTER impact", Inches(0.9), Inches(1.65), Inches(4.0), Inches(0.34), 18, NAVY, True)
    rows = [
        ("Foreign black-box solver dependency", "Inspectable indigenous solver core"),
        ("Prototype failures can be hidden by status-only checks", "Fail-loud verification contract"),
        ("Numerical scaling risk discovered late", "Threshold-gated equilibration with A/B evidence"),
        ("Roadmap guesses", "Measured limitations drive build-next priorities"),
    ]
    y = Inches(2.08)
    for before, after in rows:
        add_card(slide, "Before", before, Inches(0.9), y, Inches(4.0), Inches(0.72), LIGHT_ORANGE, rgb("F2D1BD"), ORANGE)
        arrow(slide, Inches(5.05), y + Inches(0.36), Inches(7.0), y + Inches(0.36), NAVY, 2)
        add_card(slide, "After", after, Inches(7.25), y, Inches(4.4), Inches(0.72), LIGHT_GREEN, rgb("C7E3D0"), GREEN)
        y += Inches(1.0)

    add_text(slide, "Target beneficiaries", Inches(0.95), Inches(6.25), Inches(2.4), Inches(0.26), 14, BLACK, True)
    add_text(slide, "Public-sector planners, refinery/blending schedulers, researchers, and Indian engineering teams that need auditable optimization infrastructure.", Inches(2.75), Inches(6.25), Inches(9.5), Inches(0.3), 13, DARK)


def slide_research(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "RESEARCH AND REFERENCES", 8)
    add_text(slide, "Evidence used in this deck", Inches(0.85), Inches(1.65), Inches(3.7), Inches(0.32), 18, NAVY, True)
    refs = [
        ("Project source", "README.md - prototype scope, dependency boundary, commands, measured summary."),
        ("Implementation status", "docs/status.md - implemented/tested items, known failures, roadmap."),
        ("Algorithms", "docs/algorithms.md - simplex, presolve, MILP, scaling, verification details."),
        ("Limitations", "docs/limitations.md - measured gaps: dense basis solves, B&B limits, no Netlib/MIPLIB."),
        ("Results", "docs/results.md and results/benchmark_results.csv - benchmark table and plots."),
        ("Validation run", "pytest tests/ -q; scripts/run_benchmarks.py; scripts/ab_scaling.py; scripts/plot_results.py rerun in workspace."),
        ("SIH reference", "SIH2026117.pdf - section sequence and SIH-like minimal visual language."),
    ]
    y = Inches(2.05)
    for i, (t, b) in enumerate(refs):
        col = GREEN if i in [4, 5] else NAVY
        add_card(slide, t, b, Inches(0.9 if i < 4 else 6.7), y if i < 4 else Inches(2.05 + (i - 4) * 0.95), Inches(5.25), Inches(0.72), RGBColor(255, 255, 255), rgb("D4DAE4"), col)
        if i < 4:
            y += Inches(0.95)

    add_text(slide, "Guardrails followed", Inches(6.7), Inches(5.15), Inches(2.4), Inches(0.26), 15, ORANGE, True)
    add_text(slide, "No invented team data, no commercial solver performance claim, no fake GPU result, no decorative AI imagery.", Inches(6.7), Inches(5.56), Inches(5.25), Inches(0.55), 14, DARK, True)
    add_text(slide, "Output: SIH2026119.pptx", Inches(4.0), Inches(6.65), Inches(5.3), Inches(0.35), 16, NAVY, True, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    rows = load_benchmark()

    slide_title(prs)
    slide_proposed(prs)
    slide_how(prs)
    slide_innovation(prs)
    slide_technical(prs)
    slide_feasibility(prs, rows)
    slide_impact(prs)
    slide_research(prs)

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                shape.text_frame.word_wrap = True
                shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
